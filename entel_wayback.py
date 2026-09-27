#!/usr/bin/env python3
"""
Serie de tiempo de planes móviles de entel.cl usando la Wayback Machine (web.archive.org).

Flujo:
  1. Consulta la API CDX de la Wayback Machine y obtiene la PRIMERA captura de cada mes
     (collapse=timestamp:6 => agrupa por AAAAMM y se queda con la más antigua).
  2. Descarga el HTML original de cada captura (modo "id_", sin la barra de Wayback).
  3. Extrae de forma heurística las tarjetas de planes: nombre, precio, GB, minutos, SMS,
     redes sociales libres, roaming, etc.
  4. Genera un Excel con las hojas: Planes, Resumen_Mensual y Capturas.

Uso:
  python entel_wayback.py                         # 2015-01 a hoy, URLs por defecto
  python entel_wayback.py --desde 2018 --hasta 2024
  python entel_wayback.py --url www.entel.cl/planes/ --url www.entel.cl/personas/planes-moviles/
  python entel_wayback.py --salida entel_planes.xlsx
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import statistics
import sys
import time
from dataclasses import dataclass, field, asdict
from pathlib import Path

import requests
from bs4 import BeautifulSoup

CDX_API = "https://web.archive.org/cdx/search/cdx"
SNAPSHOT_URL = "https://web.archive.org/web/{ts}id_/{url}"

# Páginas de planes que Entel ha usado a lo largo del tiempo. La portada se incluye porque
# a veces destaca los planes vigentes. Se pueden agregar/quitar con --url.
DEFAULT_URLS = [
    "www.entel.cl/planes/",
    "www.entel.cl/personas/planes/",
    "www.entel.cl/personas/planes-moviles/",
    "www.entel.cl/personas/movil/planes/",
    "www.entel.cl/planes-moviles/",
    "www.entel.cl/",
]

HEADERS = {"User-Agent": "Mozilla/5.0 (research; entel-wayback time series)"}
CACHE_DIR = Path(".cache_wayback")


# --------------------------------------------------------------------------------------
# Red
# --------------------------------------------------------------------------------------
def http_get(url: str, params: dict | None = None, retries: int = 4, pause: float = 1.5) -> requests.Response:
    last_exc: Exception | None = None
    for i in range(retries):
        try:
            r = requests.get(url, params=params, headers=HEADERS, timeout=60)
            if r.status_code == 429 or r.status_code >= 500:
                raise requests.HTTPError(f"HTTP {r.status_code}")
            return r
        except requests.RequestException as exc:
            last_exc = exc
            time.sleep(pause * (2 ** i))
    raise RuntimeError(f"No se pudo descargar {url}: {last_exc}")


def first_capture_per_month(url: str, desde: str, hasta: str) -> list[dict]:
    """Devuelve la primera captura (status 200) de cada mes para una URL."""
    params = {
        "url": url,
        "output": "json",
        "from": desde,
        "to": hasta,
        "filter": "statuscode:200",
        "collapse": "timestamp:6",  # AAAAMM -> primera captura del mes
        "fl": "timestamp,original,statuscode,mimetype,digest",
    }
    r = http_get(CDX_API, params=params)
    if not r.text.strip():
        return []
    rows = r.json()
    if not rows:
        return []
    header, data = rows[0], rows[1:]
    return [dict(zip(header, row)) for row in data]


def fetch_snapshot(ts: str, original: str) -> str:
    CACHE_DIR.mkdir(exist_ok=True)
    key = re.sub(r"[^A-Za-z0-9]+", "_", f"{ts}_{original}")[:200] + ".html"
    path = CACHE_DIR / key
    if path.exists():
        return path.read_text(encoding="utf-8", errors="replace")
    r = http_get(SNAPSHOT_URL.format(ts=ts, url=original))
    r.encoding = r.apparent_encoding or "utf-8"
    html = r.text
    path.write_text(html, encoding="utf-8")
    time.sleep(1.0)  # ser amable con archive.org
    return html


# --------------------------------------------------------------------------------------
# Extracción
# --------------------------------------------------------------------------------------
PRICE_RE = re.compile(r"\$\s?(\d{1,3}(?:[.\s]\d{3})+|\d{4,6})(?!\d)")
GB_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(GB|Gigas?|MB)\b", re.I)
UNLIMITED_DATA_RE = re.compile(r"(gigas?|datos|GB|internet)\s*(libres?|ilimitad[oa]s?)|(libres?|ilimitad[oa]s?)\s*(gigas?|datos|GB)", re.I)
MIN_RE = re.compile(r"(\d[\d.]*)\s*(?:min(?:utos)?\.?)\b", re.I)
UNLIMITED_MIN_RE = re.compile(r"min(?:utos)?\s*(libres?|ilimitad[oa]s?)|(libres?|ilimitad[oa]s?)\s*min", re.I)
SMS_RE = re.compile(r"(\d[\d.]*)\s*SMS\b", re.I)
UNLIMITED_SMS_RE = re.compile(r"SMS\s*(libres?|ilimitad[oa]s?)|(libres?|ilimitad[oa]s?)\s*SMS", re.I)
PLAN_NAME_RE = re.compile(r"\b(Plan\s+[A-Za-zÁÉÍÓÚáéíóúñÑ0-9+\-. ]{1,40}?)(?=\s*(?:\$|\d+\s*GB|Libre|Ilimitad|$|\n))", re.I)
SOCIAL = ["WhatsApp", "Facebook", "Instagram", "Twitter", "TikTok", "Waze", "Spotify", "Netflix", "YouTube", "Messenger"]
ROAMING_RE = re.compile(r"roaming", re.I)
DISCOUNT_RE = re.compile(r"(\d{1,2})\s*%\s*(?:dcto|desc(?:uento)?|off)", re.I)
MONTHS_RE = re.compile(r"(?:por|durante)\s+(\d{1,2})\s+mes", re.I)


def parse_price(s: str) -> int:
    return int(re.sub(r"\D", "", s))


def plausible_price(v: int) -> bool:
    # Planes móviles en Chile: aprox. $3.000 a $150.000 mensuales
    return 2_000 <= v <= 200_000


@dataclass
class Plan:
    nombre: str = ""
    precio_clp: int | None = None
    precio_normal_clp: int | None = None
    todos_los_precios: str = ""
    datos_gb: float | None = None
    datos_ilimitados: bool = False
    minutos: int | None = None
    minutos_ilimitados: bool = False
    sms: int | None = None
    sms_ilimitados: bool = False
    redes_sociales_libres: str = ""
    roaming: bool = False
    descuento_pct: int | None = None
    meses_promocion: int | None = None
    texto_bruto: str = ""
    metodo: str = "html"


def _normalize_ws(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def plan_from_text(text: str, metodo: str = "html") -> Plan | None:
    t = _normalize_ws(text)
    prices = [parse_price(m.group(1)) for m in PRICE_RE.finditer(t)]
    prices = [p for p in prices if plausible_price(p)]
    gb = GB_RE.search(t)
    unl_data = bool(UNLIMITED_DATA_RE.search(t))
    if not prices or not (gb or unl_data or MIN_RE.search(t)):
        return None

    p = Plan(metodo=metodo, texto_bruto=t[:500])
    p.todos_los_precios = "; ".join(f"{x:,}".replace(",", ".") for x in prices)
    # Si hay dos precios, el menor suele ser el de oferta y el mayor el normal ("antes").
    p.precio_clp = min(prices)
    if len(set(prices)) > 1:
        p.precio_normal_clp = max(prices)

    if gb:
        val = float(gb.group(1).replace(",", "."))
        p.datos_gb = val / 1024 if gb.group(2).upper() == "MB" else val
    p.datos_ilimitados = unl_data

    if UNLIMITED_MIN_RE.search(t):
        p.minutos_ilimitados = True
    else:
        m = MIN_RE.search(t)
        if m:
            p.minutos = int(re.sub(r"\D", "", m.group(1)))
    if UNLIMITED_SMS_RE.search(t):
        p.sms_ilimitados = True
    else:
        m = SMS_RE.search(t)
        if m:
            p.sms = int(re.sub(r"\D", "", m.group(1)))

    p.redes_sociales_libres = ", ".join(s for s in SOCIAL if re.search(rf"\b{s}\b", t, re.I))
    p.roaming = bool(ROAMING_RE.search(t))
    m = DISCOUNT_RE.search(t)
    if m:
        p.descuento_pct = int(m.group(1))
    m = MONTHS_RE.search(t)
    if m:
        p.meses_promocion = int(m.group(1))

    m = PLAN_NAME_RE.search(t)
    if m:
        p.nombre = _normalize_ws(m.group(1))
    elif gb:
        p.nombre = f"Plan {gb.group(1)} {gb.group(2).upper()}"
    elif unl_data:
        p.nombre = "Plan Libre"
    return p


def _card_candidates(soup: BeautifulSoup) -> list:
    """Elementos más pequeños que contienen un precio y una característica de plan."""
    candidates = []
    for el in soup.find_all(["div", "li", "article", "section", "td", "tr", "a"]):
        txt = el.get_text(" ", strip=True)
        if len(txt) > 700 or not PRICE_RE.search(txt):
            continue
        if not (GB_RE.search(txt) or UNLIMITED_DATA_RE.search(txt) or MIN_RE.search(txt)):
            continue
        candidates.append(el)
    # Quitar ancestros: si un candidato contiene a otro candidato, nos quedamos con el hijo
    cand_set = set(id(c) for c in candidates)
    result = []
    for c in candidates:
        has_child_candidate = any(id(d) in cand_set for d in c.find_all(True))
        if not has_child_candidate:
            result.append(c)
    return result


def _expand_card(el):
    """Si la tarjeta solo tiene el precio, sube 1-2 niveles para capturar nombre y detalles."""
    cur = el
    for _ in range(2):
        parent = cur.parent
        if parent is None or parent.name in ("body", "html", "[document]"):
            break
        ptxt = parent.get_text(" ", strip=True)
        if len(ptxt) > 700 or len(PRICE_RE.findall(ptxt)) > 3:
            break
        cur = parent
    return cur


def _walk_json(obj, out: list):
    if isinstance(obj, dict):
        flat = json.dumps(obj, ensure_ascii=False)
        if len(flat) < 1500 and PRICE_RE.search(flat.replace('"', " ")) is None:
            # precios en JSON muchas veces vienen sin "$": buscar claves típicas
            keys = {k.lower() for k in obj}
            if keys & {"price", "precio", "amount", "monto"} and keys & {"name", "nombre", "title", "titulo", "gb", "data", "datos"}:
                out.append(obj)
        for v in obj.values():
            _walk_json(v, out)
    elif isinstance(obj, list):
        for v in obj:
            _walk_json(v, out)


def plans_from_json_scripts(soup: BeautifulSoup) -> list[Plan]:
    """Extrae planes desde JSON embebido (Next.js __NEXT_DATA__, JSON-LD, etc.)."""
    plans = []
    for sc in soup.find_all("script"):
        raw = sc.string or ""
        if not raw or ("price" not in raw.lower() and "precio" not in raw.lower()):
            continue
        raw = raw.strip()
        try:
            data = json.loads(raw)
        except ValueError:
            continue
        objs: list = []
        _walk_json(data, objs)
        for o in objs:
            lower = {k.lower(): v for k, v in o.items()}
            price = next((lower[k] for k in ("price", "precio", "amount", "monto") if k in lower), None)
            try:
                price_int = parse_price(str(price))
            except ValueError:
                continue
            if not plausible_price(price_int):
                continue
            name = str(next((lower[k] for k in ("name", "nombre", "title", "titulo") if k in lower), ""))
            text = f"{name} ${price_int} " + " ".join(str(v) for v in o.values() if isinstance(v, (str, int, float)))
            p = plan_from_text(text, metodo="json") or Plan(metodo="json", texto_bruto=text[:500])
            p.precio_clp = price_int
            if name:
                p.nombre = name
            plans.append(p)
    return plans


def extract_plans(html: str) -> list[Plan]:
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["style", "noscript", "svg"]):
        tag.decompose()
    json_plans = plans_from_json_scripts(soup)
    for tag in soup(["script"]):
        tag.decompose()

    plans: list[Plan] = []
    seen = set()
    for el in _card_candidates(soup):
        card = _expand_card(el)
        p = plan_from_text(card.get_text(" ", strip=True))
        if not p:
            continue
        key = (p.nombre.lower(), p.precio_clp, p.datos_gb)
        if key in seen:
            continue
        seen.add(key)
        plans.append(p)

    for p in json_plans:
        key = (p.nombre.lower(), p.precio_clp, p.datos_gb)
        if key not in seen:
            seen.add(key)
            plans.append(p)
    return plans


# --------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------
PLAN_COLUMNS = [
    ("mes", "Mes (AAAA-MM)"),
    ("fecha_captura", "Fecha captura"),
    ("pagina", "Página entel.cl"),
    ("nombre", "Plan"),
    ("precio_clp", "Precio CLP"),
    ("precio_normal_clp", "Precio normal CLP"),
    ("todos_los_precios", "Todos los precios detectados"),
    ("datos_gb", "Datos (GB)"),
    ("datos_ilimitados", "Datos libres"),
    ("minutos", "Minutos"),
    ("minutos_ilimitados", "Minutos libres"),
    ("sms", "SMS"),
    ("sms_ilimitados", "SMS libres"),
    ("redes_sociales_libres", "Redes sociales libres"),
    ("roaming", "Menciona roaming"),
    ("descuento_pct", "Descuento %"),
    ("meses_promocion", "Meses promoción"),
    ("metodo", "Método extracción"),
    ("wayback_url", "URL Wayback"),
    ("texto_bruto", "Texto bruto (para revisión)"),
]


def write_excel(rows: list[dict], captures: list[dict], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, widths):
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font = header_fill, header_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # Hoja Planes
    ws = wb.active
    ws.title = "Planes"
    ws.append([h for _, h in PLAN_COLUMNS])
    for r in rows:
        ws.append([
            ("Sí" if r.get(k) is True else "No" if r.get(k) is False else r.get(k))
            for k, _ in PLAN_COLUMNS
        ])
    widths = [10, 12, 32, 26, 12, 14, 26, 10, 10, 10, 10, 8, 8, 28, 10, 10, 10, 10, 50, 80]
    style_header(ws, widths)
    for col in ("E", "F"):
        for cell in ws[col][1:]:
            cell.number_format = '"$"#,##0'
    for cell in ws["S"][1:]:
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")

    # Hoja Resumen_Mensual
    ws2 = wb.create_sheet("Resumen_Mensual")
    ws2.append(["Mes", "N° planes", "Precio mín", "Precio mediano", "Precio máx",
                "GB mín", "GB máx", "N° planes datos libres", "CLP por GB (mediana)"])
    by_month: dict[str, list[dict]] = {}
    for r in rows:
        by_month.setdefault(r["mes"], []).append(r)
    for mes in sorted(by_month):
        rs = by_month[mes]
        prices = [r["precio_clp"] for r in rs if r.get("precio_clp")]
        gbs = [r["datos_gb"] for r in rs if r.get("datos_gb")]
        clp_gb = [r["precio_clp"] / r["datos_gb"] for r in rs if r.get("precio_clp") and r.get("datos_gb")]
        ws2.append([
            mes, len(rs),
            min(prices) if prices else None,
            statistics.median(prices) if prices else None,
            max(prices) if prices else None,
            min(gbs) if gbs else None,
            max(gbs) if gbs else None,
            sum(1 for r in rs if r.get("datos_ilimitados")),
            round(statistics.median(clp_gb)) if clp_gb else None,
        ])
    style_header(ws2, [10, 10, 12, 14, 12, 8, 8, 12, 14])
    for col in ("C", "D", "E", "I"):
        for cell in ws2[col][1:]:
            cell.number_format = '"$"#,##0'

    # Hoja Capturas (log)
    ws3 = wb.create_sheet("Capturas")
    ws3.append(["Mes", "Timestamp", "Página", "N° planes extraídos", "Estado", "URL Wayback"])
    for c in captures:
        ws3.append([c["mes"], c["timestamp"], c["original"], c["n_planes"], c["estado"], c["wayback_url"]])
    style_header(ws3, [10, 16, 45, 12, 30, 70])

    # Hoja Notas
    ws4 = wb.create_sheet("Notas")
    notas = [
        "Fuente: Internet Archive Wayback Machine (web.archive.org), primera captura HTTP 200 de cada mes.",
        "Extracción heurística: revise la columna 'Texto bruto' para validar casos dudosos.",
        "Precio CLP = menor precio detectado en la tarjeta (usualmente precio oferta). Precio normal = mayor precio si hay más de uno.",
        "Páginas renderizadas solo con JavaScript pueden no traer precios en el HTML archivado; ver hoja Capturas (estado 'sin planes').",
        f"Generado: {dt.datetime.now():%Y-%m-%d %H:%M}",
    ]
    for n in notas:
        ws4.append([n])
    ws4.column_dimensions["A"].width = 120

    wb.save(path)


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------
def run(urls: list[str], desde: str, hasta: str, salida: str) -> None:
    rows: list[dict] = []
    captures: list[dict] = []
    for url in urls:
        print(f"[CDX] {url}", file=sys.stderr)
        try:
            caps = first_capture_per_month(url, desde, hasta)
        except Exception as exc:  # noqa: BLE001
            print(f"  error CDX: {exc}", file=sys.stderr)
            continue
        print(f"  {len(caps)} meses con captura", file=sys.stderr)
        for cap in caps:
            ts, original = cap["timestamp"], cap["original"]
            mes = f"{ts[:4]}-{ts[4:6]}"
            wb_url = f"https://web.archive.org/web/{ts}/{original}"
            info = {"mes": mes, "timestamp": ts, "original": original, "wayback_url": wb_url, "n_planes": 0}
            try:
                html = fetch_snapshot(ts, original)
                plans = extract_plans(html)
                info["n_planes"] = len(plans)
                info["estado"] = "ok" if plans else "sin planes detectados"
            except Exception as exc:  # noqa: BLE001
                plans = []
                info["estado"] = f"error: {exc}"[:120]
            captures.append(info)
            print(f"  {mes} {ts} -> {info['estado']} ({len(plans)})", file=sys.stderr)
            for p in plans:
                d = asdict(p)
                d.update(mes=mes, fecha_captura=f"{ts[:4]}-{ts[4:6]}-{ts[6:8]}", pagina=original, wayback_url=wb_url)
                rows.append(d)

    rows.sort(key=lambda r: (r["mes"], r["pagina"], r.get("precio_clp") or 0))
    captures.sort(key=lambda c: (c["mes"], c["original"]))
    write_excel(rows, captures, salida)
    print(f"\nListo: {len(rows)} filas de planes, {len(captures)} capturas -> {salida}", file=sys.stderr)


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--url", action="append", help="URL a consultar (repetible). Por defecto: páginas de planes de Entel.")
    ap.add_argument("--desde", default="2015", help="Año o AAAAMM inicial (default 2015)")
    ap.add_argument("--hasta", default=dt.date.today().strftime("%Y%m"), help="Año o AAAAMM final (default: mes actual)")
    ap.add_argument("--salida", default="entel_planes_moviles_wayback.xlsx", help="Archivo Excel de salida")
    args = ap.parse_args(argv)
    run(args.url or DEFAULT_URLS, args.desde, args.hasta, args.salida)


if __name__ == "__main__":
    main()
